import json
from functools import wraps
from django.conf import settings
from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
from datetime import date as date_cls, datetime
from .models import (
    Title, WorkEntry, EventEntry, PERSONS,
    LayoutDay, LayoutRace, LayoutCell, LayoutMember,
    LAYOUT_RACE_COUNT, LAYOUT_COL_COUNT, LAYOUT_DEFAULT_HEADERS, LAYOUT_COLOR_KEYS,
)

NITEI_SESSION_KEY = 'nitei_authed'


# ── nitei 専用認証デコレーター ────────────────────────
# Django ログイン済み（管理者）はそのまま通す
# nitei セッションがあればそのまま通す
# どちらでもなければ nitei ログインページへ

def nitei_login_required(view_func):
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if request.user.is_authenticated:
            return view_func(request, *args, **kwargs)
        if request.session.get(NITEI_SESSION_KEY):
            return view_func(request, *args, **kwargs)
        return redirect('nitei:login')
    return wrapper


# ── nitei ログイン / ログアウト ───────────────────────

def nitei_login(request):
    # すでに認証済みならトップへ
    if request.user.is_authenticated or request.session.get(NITEI_SESSION_KEY):
        return redirect('nitei:top')

    error = False
    if request.method == 'POST':
        pw = request.POST.get('password', '')
        if pw == getattr(settings, 'NITEI_PASSWORD', ''):
            request.session[NITEI_SESSION_KEY] = True
            request.session.set_expiry(60 * 60 * 24 * 30)  # 30日間
            return redirect('nitei:top')
        error = True

    return render(request, 'nitei/login.html', {'error': error})


def nitei_logout(request):
    request.session.pop(NITEI_SESSION_KEY, None)
    return redirect('nitei:login')


# ── ページビュー ──────────────────────────────────────

@nitei_login_required
def top(request):
    return render(request, 'nitei/top.html', {'persons': PERSONS})


@nitei_login_required
def haichi(request):
    return render(request, 'nitei/haichi.html', {
        'today':            date_cls.today().strftime('%Y-%m-%d'),
        'race_count':       LAYOUT_RACE_COUNT,
        'col_count':        LAYOUT_COL_COUNT,
        'default_headers':  json.dumps(LAYOUT_DEFAULT_HEADERS, ensure_ascii=False),
        'color_keys':       json.dumps(LAYOUT_COLOR_KEYS),
        'members':          json.dumps(_members_list(), ensure_ascii=False),
    })


@nitei_login_required
def haichi_print(request):
    """配置図の印刷ページ（最大3日分をA3横に並べ、ブラウザ印刷でPDF化）"""
    return render(request, 'nitei/haichi_print.html', {
        'dates_param':      request.GET.get('dates', ''),
        'race_count':       LAYOUT_RACE_COUNT,
        'col_count':        LAYOUT_COL_COUNT,
        'default_headers':  json.dumps(LAYOUT_DEFAULT_HEADERS, ensure_ascii=False),
    })


@nitei_login_required
def overview(request):
    return render(request, 'nitei/overview.html', {
        'persons_json': json.dumps(PERSONS, ensure_ascii=False),
    })


@nitei_login_required
def schedule_print(request):
    """個人の勤務表の印刷ページ（?person=a&sheet=N の勤務表1枚分を横長1枚に。ブラウザ印刷でPDF化）"""
    person = request.GET.get('person', '')
    if person not in PERSONS:
        return redirect('nitei:top')
    return render(request, 'nitei/schedule_print.html', {
        'person':      person,
        'person_name': PERSONS[person],
        'sheet_param': request.GET.get('sheet', ''),
    })


@nitei_login_required
def overview_print(request):
    """全員一覧の印刷ページ（?month=YYYY-MM の1か月分をA3横1枚に。ブラウザ印刷でPDF化）"""
    return render(request, 'nitei/overview_print.html', {
        'persons_json': json.dumps(PERSONS, ensure_ascii=False),
        'month_param':  request.GET.get('month', ''),
    })


def api_overview(request):
    if not (request.user.is_authenticated or request.session.get(NITEI_SESSION_KEY)):
        return JsonResponse({'error': 'unauthorized'}, status=403)
    sheet_index   = int(request.GET.get('sheet_index', 0))
    section_index = int(request.GET.get('section_index', 0))

    data = {}
    for person in PERSONS:
        pdata = {}
        for e in WorkEntry.objects.filter(
                person=person, sheet_index=sheet_index, section_index=section_index):
            pdata[f"w_{e.day_index}_{e.row_type}"] = e.status
        for e in EventEntry.objects.filter(
                person=person, sheet_index=sheet_index, section_index=section_index):
            pdata[f"e_{e.day_index}"] = e.time_text
        data[person] = pdata

    titles = list(Title.objects.values('id', 'date_from', 'date_to', 'venue', 'title'))
    for t in titles:
        t['date_from'] = t['date_from'].strftime('%Y/%m/%d')
        t['date_to']   = t['date_to'].strftime('%Y/%m/%d')

    return JsonResponse({'data': data, 'titles': titles})


def schedule(request, person):
    if person not in PERSONS:
        return redirect('nitei:top')

    if person == 'a':
        if not request.user.is_authenticated:
            return redirect(f'/accounts/login/?next={request.path}')
    else:
        if not request.user.is_authenticated and not request.session.get(NITEI_SESSION_KEY):
            return redirect('nitei:login')

    return render(request, 'nitei/index.html', {
        'person':      person,
        'person_name': PERSONS[person],
        'persons':     PERSONS,
    })


# ── 開催タイトル API ──────────────────────────────────

@nitei_login_required
def api_titles(request):
    titles = list(Title.objects.values('id', 'date_from', 'date_to', 'venue', 'title'))
    for t in titles:
        t['date_from'] = t['date_from'].strftime('%Y/%m/%d')
        t['date_to']   = t['date_to'].strftime('%Y/%m/%d')
    return JsonResponse(titles, safe=False)


# ── 勤務記録 API ──────────────────────────────────────

@nitei_login_required
def api_schedule(request):
    person = request.GET.get('person', 'a')
    if person not in PERSONS:
        return JsonResponse({'error': 'invalid person'}, status=400)
    entries = WorkEntry.objects.filter(person=person)
    data = {f"w_{e.sheet_index}_{e.section_index}_{e.day_index}_{e.row_type}": e.status
            for e in entries}
    return JsonResponse(data)


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_schedule_save(request):
    try:
        body   = json.loads(request.body)
        key    = body['key']
        status = body['status']
        person = body.get('person', 'a')
    except (KeyError, json.JSONDecodeError):
        return JsonResponse({'error': 'invalid'}, status=400)

    if person not in PERSONS:
        return JsonResponse({'error': 'invalid person'}, status=400)

    parts = key.split('_')
    if len(parts) != 5 or parts[0] != 'w':
        return JsonResponse({'error': 'bad key'}, status=400)

    _, si, sec, di, rt = parts
    if status == '':
        WorkEntry.objects.filter(
            person=person,
            sheet_index=int(si), section_index=int(sec), day_index=int(di), row_type=int(rt)
        ).delete()
    else:
        WorkEntry.objects.update_or_create(
            person=person,
            sheet_index=int(si), section_index=int(sec), day_index=int(di), row_type=int(rt),
            defaults={'status': status}
        )
    return JsonResponse({'ok': True})


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_schedule_clear(request):
    try:
        body        = json.loads(request.body)
        person      = body.get('person', 'a')
        sheet_index = int(body['sheet_index'])
    except (KeyError, ValueError, json.JSONDecodeError):
        return JsonResponse({'error': 'invalid'}, status=400)

    if person not in PERSONS:
        return JsonResponse({'error': 'invalid person'}, status=400)

    WorkEntry.objects.filter(person=person, sheet_index=sheet_index).delete()
    return JsonResponse({'ok': True})


# ── 開催行 時間メモ API ────────────────────────────────

@nitei_login_required
def api_events(request):
    person = request.GET.get('person', 'a')
    if person not in PERSONS:
        return JsonResponse({'error': 'invalid person'}, status=400)
    entries = EventEntry.objects.filter(person=person)
    data = {f"e_{e.sheet_index}_{e.section_index}_{e.day_index}": e.time_text
            for e in entries}
    return JsonResponse(data)


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_events_save(request):
    try:
        body      = json.loads(request.body)
        key       = body['key']
        time_text = body['time_text']
        person    = body.get('person', 'a')
    except (KeyError, json.JSONDecodeError):
        return JsonResponse({'error': 'invalid'}, status=400)

    if person not in PERSONS:
        return JsonResponse({'error': 'invalid person'}, status=400)

    parts = key.split('_')
    if len(parts) != 4 or parts[0] != 'e':
        return JsonResponse({'error': 'bad key'}, status=400)

    _, si, sec, di = parts
    if time_text == '':
        EventEntry.objects.filter(
            person=person,
            sheet_index=int(si), section_index=int(sec), day_index=int(di)
        ).delete()
    else:
        EventEntry.objects.update_or_create(
            person=person,
            sheet_index=int(si), section_index=int(sec), day_index=int(di),
            defaults={'time_text': time_text}
        )
    return JsonResponse({'ok': True})


# ── 配置図 API ────────────────────────────────────────

def _parse_date(value):
    """'YYYY-MM-DD' を date に。不正なら None"""
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None


def _clean_time(value):
    """HH:MM 形式だけ通す。それ以外は空文字"""
    if not isinstance(value, str):
        return ''
    value = value.strip()
    if not value:
        return ''
    try:
        datetime.strptime(value, '%H:%M')
    except ValueError:
        return ''
    return value


@nitei_login_required
def api_layout(request):
    day_date = _parse_date(request.GET.get('date'))
    if day_date is None:
        return JsonResponse({'error': 'invalid date'}, status=400)

    day = LayoutDay.objects.filter(date=day_date).first()
    if day is None:
        return JsonResponse({
            'date':    day_date.strftime('%Y-%m-%d'),
            'headers': LAYOUT_DEFAULT_HEADERS,
            'races':   [],
            'cells':   {},
            'colors':  {},
            'exists':  False,
        })

    headers = day.headers or LAYOUT_DEFAULT_HEADERS
    headers = (list(headers) + [''] * LAYOUT_COL_COUNT)[:LAYOUT_COL_COUNT]

    cells  = {}
    colors = {}
    for c in day.cells.all():
        key = f"{c.race}_{c.col}"
        if c.text:
            cells[key] = c.text
        if c.color:
            colors[key] = c.color

    return JsonResponse({
        'date':    day_date.strftime('%Y-%m-%d'),
        'headers': headers,
        'races': [{
            'race':      r.race,
            'start':     r.start_time,
            'close':     r.close_time,
            'highlight': r.highlight,
        } for r in day.races.all()],
        'cells':  cells,
        'colors': colors,
        'exists': True,
    })


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_layout_save(request):
    """1日分をまとめて保存（差し替え）"""
    try:
        body = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'error': 'invalid json'}, status=400)

    day_date = _parse_date(body.get('date'))
    if day_date is None:
        return JsonResponse({'error': 'invalid date'}, status=400)

    headers = body.get('headers') or []
    if not isinstance(headers, list):
        return JsonResponse({'error': 'invalid headers'}, status=400)
    headers = [str(h)[:50] for h in headers[:LAYOUT_COL_COUNT]]
    headers += [''] * (LAYOUT_COL_COUNT - len(headers))

    day, _ = LayoutDay.objects.update_or_create(
        date=day_date, defaults={'headers': headers})

    # レース時刻
    day.races.all().delete()
    races = []
    for item in (body.get('races') or []):
        try:
            race = int(item.get('race'))
        except (TypeError, ValueError):
            continue
        if not 1 <= race <= LAYOUT_RACE_COUNT:
            continue
        start     = _clean_time(item.get('start'))
        close     = _clean_time(item.get('close'))
        highlight = bool(item.get('highlight'))
        if not start and not close and not highlight:
            continue
        races.append(LayoutRace(day=day, race=race, start_time=start,
                                close_time=close, highlight=highlight))
    LayoutRace.objects.bulk_create(races)

    # 配置セル（テキストと背景色。どちらか一方だけでも保存する）
    def parse_key(key):
        parts = str(key).split('_')
        if len(parts) != 2:
            return None
        try:
            race, col = int(parts[0]), int(parts[1])
        except ValueError:
            return None
        if not (1 <= race <= LAYOUT_RACE_COUNT and 0 <= col < LAYOUT_COL_COUNT):
            return None
        return (race, col)

    merged = {}
    for key, text in (body.get('cells') or {}).items():
        rc = parse_key(key)
        if rc is None:
            continue
        text = str(text or '').strip()[:200]
        if text:
            merged.setdefault(rc, {'text': '', 'color': ''})['text'] = text

    for key, color in (body.get('colors') or {}).items():
        rc = parse_key(key)
        if rc is None:
            continue
        color = str(color or '').strip()
        if color in LAYOUT_COLOR_KEYS:
            merged.setdefault(rc, {'text': '', 'color': ''})['color'] = color

    day.cells.all().delete()
    LayoutCell.objects.bulk_create([
        LayoutCell(day=day, race=rc[0], col=rc[1], text=v['text'], color=v['color'])
        for rc, v in merged.items()
    ])

    return JsonResponse({'ok': True})


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_layout_clear(request):
    try:
        body = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'error': 'invalid json'}, status=400)

    day_date = _parse_date(body.get('date'))
    if day_date is None:
        return JsonResponse({'error': 'invalid date'}, status=400)

    LayoutDay.objects.filter(date=day_date).delete()
    return JsonResponse({'ok': True})


# ── 配置図: 発売時刻の Excel 取り込み ─────────────────
# テンプレートは、人やAIが元資料（PDF・別のExcel等）から書き写すことを前提に、
# 入力セル（黄色）・行とレースの対応・書式・してはいけないことを「説明」シートに明記する。
# 取り込みは「発売時刻」シートの「A列にレース番号がある行」と「対象日」の行だけを読む
EXCEL_MAX_BYTES  = 2 * 1024 * 1024
EXCEL_SHEET      = '発売時刻'
EXCEL_FIRST_ROW  = 6          # 1R の行（レース番号 + 5）
EXCEL_DATE_LABEL = '対象日'

EXCEL_GUIDE = [
    ('このファイルの用途',
     '業務ツール「日程 > 配置図」の「📥 時刻Excel読込」で読み込み、画面に表示中の日付の'
     '各レースの「発売開始」と「締め切り」の時刻を一括で入力するためのファイルです。'
     '人またはAIが、出走表などの元資料（PDF・Excel・画像など）から時刻を書き写して使います。'),
    ('入力するセル',
     'シート「発売時刻」の黄色のセルだけです。\n'
     '・B6〜B17：1R〜12R の発売開始\n'
     '・C6〜C17：1R〜12R の締め切り\n'
     '・B3：対象日（任意）\n'
     'これ以外のセルには書き込まないでください。'),
    ('行とレースの対応',
     '6行目=1R、7行目=2R、8行目=3R … 17行目=12R（行番号 = レース番号 + 5）。\n'
     'A列にレース番号（1R〜12R）が入っているので、同じ行の B列・C列 に書きます。'),
    ('発売開始（B列）',
     'そのレースの投票券の発売が始まる時刻。\n'
     '元資料での呼び名の例：「発売開始」「発売開始時刻」「発売開始予定」。'),
    ('締め切り（C列）',
     'そのレースの投票の締め切り時刻。\n'
     '元資料での呼び名の例：「締切」「締切予定時刻」「投票締切」「発売締切」「締切時刻」。\n'
     '「発走時刻」「スタート時刻」「展示時刻」は締め切りとは別の時刻なので入れないでください。'
     '締め切りが資料に無い場合は、推測で埋めず空欄にしてください。'),
    ('時刻の書き方',
     '24時間表記の「時:分」（例：10:30、20:45）。Excelの時刻値でも、文字列 "10:30" でも構いません。\n'
     '秒は不要です。「午前／午後」「AM／PM」は使わないでください。'),
    ('日をまたぐ場合',
     'ナイター等で締め切りが翌日の0時台になる場合も、日付を付けずに 0:15 のように書いてください。'
     '発売時間（D列）は日またぎを考慮して自動計算されます。'),
    ('分からない・無いレース',
     '空欄のままにしてください（推測で埋めない）。空欄のセルは取り込まれず、画面に入力済みの時刻がそのまま残ります。'
     '12レースより少ない開催日は、残りの行を空欄にします。'),
    ('対象日（B3）',
     '元資料の開催日を日付で書きます（例：2026/10/17）。任意です。'
     '書いておくと、画面で表示中の日付と違うときに読み込み前に確認が出ます。'),
    ('変更しないもの',
     'シート名「発売時刻」、A列のレース番号、5行目の見出し、行と列の位置。'
     '行や列の挿入・削除、セルの結合はしないでください。'),
    ('D列・E列',
     'D列「発売時間」は計算式なので書き込み不要です（取り込まれません）。'
     'E列「メモ」は自由記入欄で、取り込まれません（資料の出典などを書いておけます）。'),
    ('保存形式',
     '.xlsx のまま保存してください（CSV・PDF・画像では読み込めません）。ファイル名は自由です。'),
    ('記入例',
     '1R の発売開始が 10:05、締め切りが 10:28 なら → B6 に 10:05、C6 に 10:28。\n'
     '2R の発売開始が 10:33、締め切りが 10:56 なら → B7 に 10:33、C7 に 10:56。'),
    ('仕上げの確認',
     '① B3 と B6〜C17 以外に書き込んでいない\n'
     '② 時刻はすべて「時:分」の形になっている\n'
     '③ 各行で 発売開始 → 締め切り の順になっている（日またぎを除く）\n'
     '④ レース番号が進むほど時刻も遅くなっている\n'
     '⑤ 資料から読み取れなかった値は空欄のままにした'),
]


@nitei_login_required
def layout_excel_template(request):
    """発売時刻取り込み用の Excel テンプレートをダウンロード（入力シート＋説明シート）"""
    from io import BytesIO
    from urllib.parse import quote
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    thin   = Side(style='thin', color='888888')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')
    wrap   = Alignment(wrap_text=True, vertical='top')
    head   = PatternFill('solid', fgColor='DDE6F0')
    entry  = PatternFill('solid', fgColor='FFF2A8')   # 入力セル（黄色）
    fixed  = PatternFill('solid', fgColor='EEEEEE')   # 変更しないセル
    last   = EXCEL_FIRST_ROW + LAYOUT_RACE_COUNT - 1

    wb = Workbook()
    ws = wb.active
    ws.title = EXCEL_SHEET

    ws['A1'] = '配置図 発売時刻 取り込みシート'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = (f'入力するのは黄色のセルだけです（B3 と B{EXCEL_FIRST_ROW}:C{last}）。'
                '書き方のルールは「説明」シートを参照してください。')
    ws['A2'].font = Font(color='C00000', bold=True)
    ws.merge_cells('A2:E2')

    ws['A3'] = EXCEL_DATE_LABEL
    ws['A3'].font = Font(bold=True)
    ws['A3'].fill = fixed
    ws['A3'].alignment = center
    ws['A3'].border = border
    ws['B3'].fill = entry
    ws['B3'].border = border
    ws['B3'].number_format = 'yyyy/mm/dd'
    ws['B3'].alignment = center
    ws['C3'] = '← 任意。元資料の開催日（例 2026/10/17）。画面の日付と違うと読み込み時に確認が出ます'
    ws['C3'].font = Font(size=9, color='555555')
    ws.merge_cells('C3:E3')

    for col, title in enumerate(['レース', '発売開始', '締め切り', '発売時間（自動計算）', 'メモ（取り込まれない）'], start=1):
        c = ws.cell(row=EXCEL_FIRST_ROW - 1, column=col, value=title)
        c.font = Font(bold=True)
        c.fill = head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border

    times_dv = DataValidation(
        type='time', operator='between', formula1='0', formula2='0.999305555555556',
        allow_blank=True, showErrorMessage=True,
        errorTitle='時刻の形式', error='24時間表記の「時:分」（例 10:30）で入力してください')
    ws.add_data_validation(times_dv)

    for r in range(1, LAYOUT_RACE_COUNT + 1):
        row = EXCEL_FIRST_ROW + r - 1
        a = ws.cell(row=row, column=1, value=f'{r}R')
        a.fill, a.font = fixed, Font(bold=True)
        for col in range(1, 6):
            c = ws.cell(row=row, column=col)
            c.border = border
            c.alignment = center
        for col in (2, 3):
            c = ws.cell(row=row, column=col)
            c.fill = entry
            c.number_format = 'hh:mm'
        d = ws.cell(row=row, column=4, value=f'=IF(OR(B{row}="",C{row}=""),"",MOD(C{row}-B{row},1))')
        d.number_format = 'hh:mm'
        d.fill = fixed
    times_dv.add(f'B{EXCEL_FIRST_ROW}:C{last}')

    for col, width in zip('ABCDE', (10, 14, 14, 16, 30)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[EXCEL_FIRST_ROW - 1].height = 30
    ws.freeze_panes = f'A{EXCEL_FIRST_ROW}'

    guide = wb.create_sheet('説明')
    guide['A1'] = '「発売時刻」シートの書き方（このファイルを編集する人・AI向け）'
    guide['A1'].font = Font(bold=True, size=13)
    for col, title in enumerate(['項目', '内容'], start=1):
        c = guide.cell(row=3, column=col, value=title)
        c.font, c.fill, c.border, c.alignment = Font(bold=True), head, border, center
    for i, (item, text) in enumerate(EXCEL_GUIDE, start=4):
        a = guide.cell(row=i, column=1, value=item)
        a.font, a.border, a.alignment = Font(bold=True), border, wrap
        b = guide.cell(row=i, column=2, value=text)
        b.border, b.alignment = border, wrap
        guide.row_dimensions[i].height = 15 * (text.count('\n') + 1 + len(text) // 55)
    guide.column_dimensions['A'].width = 22
    guide.column_dimensions['B'].width = 90

    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = '配置図_発売時刻テンプレート.xlsx'
    resp['Content-Disposition'] = (
        f"attachment; filename=\"haichi_times_template.xlsx\"; filename*=UTF-8''{quote(filename)}")
    return resp


def _excel_race(value):
    """A列の値をレース番号（1〜12）に。'1R' / '１Ｒ' / 1 などを許容。該当しなければ None"""
    import re
    import unicodedata
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        n = int(value) if float(value).is_integer() else None
    else:
        m = re.fullmatch(r'(\d{1,2})\s*R?', unicodedata.normalize('NFKC', str(value)).strip(), re.I)
        n = int(m.group(1)) if m else None
    return n if n is not None and 1 <= n <= LAYOUT_RACE_COUNT else None


def _excel_time(value):
    """セルの値を 'HH:MM' に。空欄は ''、時刻として読めなければ None"""
    import re
    import unicodedata
    from datetime import time as time_cls, timedelta

    def fmt(h, m):
        return f'{h:02d}:{m:02d}' if 0 <= h <= 23 and 0 <= m <= 59 else None

    def from_clock(t):
        # Excel の時刻は浮動小数なので 10:29:59.99 のようになる。分に四捨五入する
        mins = round((t.hour * 3600 + t.minute * 60 + t.second + t.microsecond / 1e6) / 60)
        return fmt(mins // 60, mins % 60) if mins < 24 * 60 else None

    if value is None or (isinstance(value, str) and not value.strip()):
        return ''
    if isinstance(value, datetime):
        # 時刻書式のセルに 1135 のように数字だけ入れると「1135日後の0:00」になる。
        # その場合は元の数字を HHMM として読む
        serial = (value - datetime(1899, 12, 30)).days
        if serial >= 1 and value.hour == value.minute == value.second == 0:
            return fmt(serial // 100, serial % 100)
        return from_clock(value)
    if isinstance(value, time_cls):
        return from_clock(value)
    if isinstance(value, timedelta):
        mins = round(value.total_seconds() / 60)
        return fmt(mins // 60, mins % 60)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 0 <= value < 1:                       # Excel の時刻（1日=1.0）
            mins = round(value * 24 * 60)
            return fmt(mins // 60, mins % 60) if mins < 24 * 60 else None
        if float(value).is_integer():            # 1030 → 10:30
            n = int(value)
            return fmt(n // 100, n % 100)
        return None
    s = unicodedata.normalize('NFKC', str(value)).strip()
    m = (re.fullmatch(r'(\d{1,2}):(\d{1,2})(?::\d{1,2})?', s)
         or re.fullmatch(r'(\d{1,2})時(\d{1,2})分?', s)
         or re.fullmatch(r'(\d{1,2})(\d{2})', s))
    return fmt(int(m.group(1)), int(m.group(2))) if m else None


def _excel_date(value):
    """対象日のセルを 'YYYY-MM-DD' に。空欄・読めない場合は ''"""
    import re
    import unicodedata
    from datetime import date as date_only
    if isinstance(value, (datetime, date_only)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, str):
        m = re.fullmatch(r'(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})日?',
                         unicodedata.normalize('NFKC', value).strip())
        if m:
            try:
                return date_only(int(m.group(1)), int(m.group(2)), int(m.group(3))).strftime('%Y-%m-%d')
            except ValueError:
                return ''
    return ''


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def layout_excel_import(request):
    """Excel から発売開始・締め切り（と対象日）を読み取って返す（保存は画面側が通常の保存で行う）"""
    from openpyxl import load_workbook

    f = request.FILES.get('file')
    if f is None:
        return JsonResponse({'error': 'ファイルが選ばれていません'}, status=400)
    if f.size > EXCEL_MAX_BYTES:
        return JsonResponse({'error': 'ファイルが大きすぎます（2MBまで）'}, status=400)
    try:
        wb = load_workbook(f, data_only=True, read_only=True)
    except Exception:
        return JsonResponse({'error': 'Excelファイル（.xlsx）として読み込めませんでした'}, status=400)

    races, warnings, target_date = {}, [], ''
    try:
        ws = wb[EXCEL_SHEET] if EXCEL_SHEET in wb.sheetnames else wb.worksheets[0]
        for row in ws.iter_rows(max_col=3, values_only=True):
            if not row:
                continue
            if isinstance(row[0], str) and row[0].strip() == EXCEL_DATE_LABEL:
                target_date = _excel_date(row[1] if len(row) > 1 else None)
                continue
            race = _excel_race(row[0])
            if race is None:
                continue
            times = {}
            for label, key, raw in (('発売開始', 'start', row[1] if len(row) > 1 else None),
                                    ('締め切り', 'close', row[2] if len(row) > 2 else None)):
                t = _excel_time(raw)
                if t is None:
                    warnings.append(f'{race}R {label}「{raw}」')
                elif t:
                    times[key] = t
            if times:
                races[race] = {'start': times.get('start', ''), 'close': times.get('close', '')}
    finally:
        wb.close()

    return JsonResponse({'races': races, 'warnings': warnings, 'date': target_date})


# ── 配置図の名簿（クリック循環の順序・自動着色の色） ────

def _members_list():
    return [{'id': m.id, 'name': m.name, 'color': m.color, 'order': m.order}
            for m in LayoutMember.objects.all()]


def _members_response():
    return JsonResponse({'members': _members_list()})


@nitei_login_required
def api_members(request):
    return _members_response()


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_member_save(request):
    """名簿の追加 or 更新（id あり=更新）。payload: {id?, name?, color?}
    改名時は、配置済みセル（全日付）の氏名・色も追従して書き換える"""
    try:
        body = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'error': 'invalid json'}, status=400)

    name  = str(body.get('name') or '').strip()[:20]
    color = str(body.get('color') or '').strip()

    member_id = body.get('id')
    if member_id:
        m = LayoutMember.objects.filter(id=member_id).first()
        if not m:
            return JsonResponse({'error': 'not found'}, status=404)
        if name and name != m.name:
            if LayoutMember.objects.filter(name=name).exclude(id=m.id).exists():
                return JsonResponse({'error': f'「{name}」は登録済みです'}, status=400)
            old = m.name
            m.name = name
            # 改名を配置済みセルへ反映（過去日も含む。同一人物の表記修正を想定）
            LayoutCell.objects.filter(text=old).update(text=name)
        if color in LAYOUT_COLOR_KEYS:
            m.color = color
            # 色変更も配置済みセルへ反映
            LayoutCell.objects.filter(text=m.name).update(color=color)
        m.save()
        return _members_response()

    # 新規追加
    if not name:
        return JsonResponse({'error': '氏名を入力してください'}, status=400)
    if LayoutMember.objects.filter(name=name).exists():
        return JsonResponse({'error': f'「{name}」は登録済みです'}, status=400)
    if color not in LAYOUT_COLOR_KEYS:
        # 未指定なら、使用数が最も少ない色を自動で割り当てる
        used = {c: 0 for c in LAYOUT_COLOR_KEYS}
        for m in LayoutMember.objects.all():
            if m.color in used:
                used[m.color] += 1
        color = min(LAYOUT_COLOR_KEYS, key=lambda c: used[c])
    last = LayoutMember.objects.order_by('-order').first()
    LayoutMember.objects.create(
        name=name, color=color, order=(last.order + 1) if last else 0)
    return _members_response()


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_member_delete(request):
    """名簿から削除。配置済みセルの文字は残す（履歴保全。色と循環からは外れる）"""
    try:
        body = json.loads(request.body)
        member_id = int(body['id'])
    except (KeyError, ValueError, json.JSONDecodeError):
        return JsonResponse({'error': 'invalid'}, status=400)
    LayoutMember.objects.filter(id=member_id).delete()
    return _members_response()


@nitei_login_required
@csrf_exempt
@require_http_methods(['POST'])
def api_member_reorder(request):
    """名簿の並び順（=クリック循環の順序）を更新。payload: {ids: [...]}"""
    try:
        body = json.loads(request.body)
        ids = [int(x) for x in body['ids']]
    except (KeyError, TypeError, ValueError, json.JSONDecodeError):
        return JsonResponse({'error': 'invalid'}, status=400)
    owned = {m.id: m for m in LayoutMember.objects.filter(id__in=ids)}
    order = 0
    for mid in ids:
        m = owned.get(mid)
        if m:
            m.order = order
            m.save(update_fields=['order'])
            order += 1
    return _members_response()
